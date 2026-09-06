from __future__ import annotations

import copy
import csv
import io
import json
from collections.abc import Iterable
from typing import Any

from refaat.core.text import normalize_text
from refaat.domain.models import ExtractedDocument, JsonObject
from refaat.errors import ValidationError
from refaat.ports import DocumentExtractor


class TextExtractor:
    media_types = frozenset(
        {
            "text/plain",
            "text/markdown",
            "text/x-markdown",
            "application/xml",
            "text/xml",
            "application/yaml",
            "text/yaml",
            "application/toml",
            "application/sql",
            "application/graphql",
            "text/x-python",
            "text/javascript",
            "text/typescript",
            "text/jsx",
            "text/tsx",
            "text/x-java-source",
            "text/x-go",
            "text/x-rust",
            "text/x-php",
            "text/x-ruby",
            "text/x-shellscript",
            "text/x-sql",
            "text/x-c",
            "text/x-c++src",
            "text/css",
        }
    )

    def extract(self, content: bytes, *, metadata: JsonObject) -> ExtractedDocument:
        encoding = str(metadata.get("encoding", "utf-8"))
        try:
            text = content.decode(encoding)
        except (UnicodeDecodeError, LookupError) as exc:
            raise ValidationError(f"Document cannot be decoded as {encoding}") from exc
        return ExtractedDocument(text=text, metadata={"encoding": encoding})


class JsonExtractor:
    media_types = frozenset({"application/json", "application/ld+json"})

    def extract(self, content: bytes, *, metadata: JsonObject) -> ExtractedDocument:
        try:
            value = json.loads(content)
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValidationError("Invalid JSON document") from exc
        return ExtractedDocument(text=self._flatten(value), metadata={"format": "json"})

    def _flatten(self, value: Any, path: str = "$") -> str:
        lines: list[str] = []
        if isinstance(value, dict):
            for key, child in value.items():
                lines.append(self._flatten(child, f"{path}.{key}"))
        elif isinstance(value, list):
            for index, child in enumerate(value):
                lines.append(self._flatten(child, f"{path}[{index}]"))
        else:
            lines.append(f"{path}: {value}")
        return "\n".join(line for line in lines if line)


class CsvExtractor:
    media_types = frozenset({"text/csv", "application/csv"})

    def extract(self, content: bytes, *, metadata: JsonObject) -> ExtractedDocument:
        try:
            text = content.decode(str(metadata.get("encoding", "utf-8-sig")))
        except UnicodeDecodeError as exc:
            raise ValidationError("CSV document cannot be decoded") from exc
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        lines = [
            "; ".join(f"{key}: {value}" for key, value in row.items() if value not in {None, ""}) for row in reader
        ]
        return ExtractedDocument(text="\n".join(lines), metadata={"format": "csv"})


class HtmlExtractor:
    media_types = frozenset({"text/html", "application/xhtml+xml"})

    def extract(self, content: bytes, *, metadata: JsonObject) -> ExtractedDocument:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise RuntimeError("Install refaat[documents] for HTML extraction") from exc
        soup = BeautifulSoup(content, "html.parser")
        for element in soup(["script", "style", "noscript", "template"]):
            element.decompose()
        title = soup.title.string.strip() if soup.title and soup.title.string else None
        return ExtractedDocument(
            text=normalize_text(soup.get_text("\n")),
            metadata={"title": title} if title else {},
        )


class PdfExtractor:
    media_types = frozenset({"application/pdf"})

    def extract(self, content: bytes, *, metadata: JsonObject) -> ExtractedDocument:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Install refaat[documents] for PDF extraction") from exc
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            password = metadata.get("password")
            if not password or reader.decrypt(str(password)) == 0:
                raise ValidationError("Encrypted PDF requires a valid password")
        pages = [page.extract_text() or "" for page in reader.pages]
        return ExtractedDocument(
            text="\n\n".join(f"[Page {index + 1}]\n{text}" for index, text in enumerate(pages)),
            metadata={"pages": len(pages)},
        )


class DocxExtractor:
    media_types = frozenset({"application/vnd.openxmlformats-officedocument.wordprocessingml.document"})

    def extract(self, content: bytes, *, metadata: JsonObject) -> ExtractedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install refaat[documents] for DOCX extraction") from exc
        document = Document(io.BytesIO(content))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return ExtractedDocument(text="\n".join(lines), metadata={"format": "docx"})


class ExtractorRegistry:
    def __init__(self, extractors: Iterable[DocumentExtractor] | None = None) -> None:
        self._extractors: dict[str, DocumentExtractor] = {}
        for extractor in extractors or self.default_extractors():
            self.register(extractor)

    @staticmethod
    def default_extractors() -> tuple[DocumentExtractor, ...]:
        return (TextExtractor(), JsonExtractor(), CsvExtractor(), HtmlExtractor(), PdfExtractor(), DocxExtractor())

    def register(self, extractor: DocumentExtractor) -> None:
        for media_type in extractor.media_types:
            if media_type in self._extractors:
                raise ValueError(f"Extractor already registered for {media_type}")
            self._extractors[media_type] = extractor

    def extract(self, media_type: str, content: bytes, metadata: JsonObject) -> ExtractedDocument:
        canonical = media_type.split(";", maxsplit=1)[0].strip().lower()
        extractor = self._extractors.get(canonical)
        if extractor is None:
            raise ValidationError(f"Unsupported media type: {canonical}")
        # An isolated copy, so a plugin extractor cannot mutate the caller's metadata in place and inject a
        # governance key into the merge base (the output governance keys are stripped separately).
        result = extractor.extract(content, metadata=copy.deepcopy(metadata))
        if not result.text.strip():
            raise ValidationError("Extractor produced no text")
        return result

    @property
    def media_types(self) -> frozenset[str]:
        return frozenset(self._extractors)
