from __future__ import annotations

import io

import pytest
from docx import Document

from refaat.application.runtime import Refaat
from refaat.domain.models import DocumentInput

pytestmark = pytest.mark.integration

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _make_docx() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Release runbook")
    document.add_paragraph("A production release requires peer review and a rollback plan before deployment approval.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "owner"
    table.cell(0, 1).text = "platform team"
    document.save(buffer)
    return buffer.getvalue()


def _ingest_and_cite(runtime: Refaat, *, source_id: str, media_type: str, content: bytes, query: str) -> None:
    project = runtime.project(workspace="formats", corpus=source_id)
    result = runtime.ingest(
        project.workspace_id,
        project.corpus_id,
        DocumentInput(
            source_id=source_id,
            title=source_id,
            media_type=media_type,
            content=content,
        ),
    )
    assert result.chunks >= 1

    stored = {
        chunk.id
        for document in project.list_documents()
        for chunk in runtime.store.list_document_chunks(project.workspace_id, project.corpus_id, str(document["id"]))
    }
    assert stored

    response = project.ask(query)
    assert response.citations
    assert all(citation.chunk_id in stored for citation in response.citations)


def test_docx_ingests_into_citable_chunks(runtime: Refaat) -> None:
    _ingest_and_cite(
        runtime,
        source_id="runbook.docx",
        media_type=DOCX_MEDIA_TYPE,
        content=_make_docx(),
        query="What does a production release require before deployment?",
    )


def test_html_ingests_into_citable_chunks(runtime: Refaat) -> None:
    html = (
        b"<html><head><title>Handbook</title><script>bad()</script></head>"
        b"<body><h1>Deployment policy</h1><p>Every production release requires an approved rollback plan.</p>"
        b"</body></html>"
    )
    _ingest_and_cite(
        runtime,
        source_id="handbook.html",
        media_type="text/html",
        content=html,
        query="What is required for a production release?",
    )


def test_csv_ingests_into_citable_chunks(runtime: Refaat) -> None:
    csv_content = (
        b"control,owner\npeer review,release manager\nrollback plan,platform team\nthreat review,security team\n"
    )
    _ingest_and_cite(
        runtime,
        source_id="controls.csv",
        media_type="text/csv",
        content=csv_content,
        query="Who owns the rollback plan control?",
    )


def test_docx_extractor_reports_actionable_error_when_dependency_absent(
    runtime: Refaat, monkeypatch: pytest.MonkeyPatch
) -> None:
    real_import = __import__

    def blocked_import(name: str, *args: object, **kwargs: object) -> object:
        if name == "docx" or name.startswith("docx."):
            raise ImportError(name)
        return real_import(name, *args, **kwargs)

    content = _make_docx()
    project = runtime.project(workspace="formats", corpus="docx-missing")
    document = DocumentInput(
        source_id="runbook.docx",
        title="Runbook",
        media_type=DOCX_MEDIA_TYPE,
        content=content,
    )
    monkeypatch.setattr("builtins.__import__", blocked_import)
    with pytest.raises(RuntimeError, match="refaat\\[documents\\] for DOCX"):
        runtime.ingest(project.workspace_id, project.corpus_id, document)
