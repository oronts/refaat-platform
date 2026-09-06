from __future__ import annotations

import io

import pytest
from docx import Document
from reportlab.pdfgen.canvas import Canvas

from refaat.errors import ValidationError
from refaat.infrastructure.extractors import ExtractorRegistry, TextExtractor


def make_docx() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Runbook")
    document.add_paragraph("Restart only after approval.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "owner"
    table.cell(0, 1).text = "platform"
    document.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    buffer = io.BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(72, 720, "Production evidence requires citations.")
    canvas.save()
    return buffer.getvalue()


def test_builtin_extractors_cover_structured_and_document_formats() -> None:
    registry = ExtractorRegistry()
    assert "$.service.name: refaat" in registry.extract("application/json", b'{"service":{"name":"refaat"}}', {}).text
    assert "name: Refaat" in registry.extract("text/csv", b"name,kind\nRefaat,runtime\n", {}).text
    html = registry.extract(
        "text/html",
        b"<html><head><title>Evidence</title><script>bad()</script></head><body><h1>Safe context</h1></body></html>",
        {},
    )
    assert html.metadata["title"] == "Evidence"
    assert "bad" not in html.text
    assert "Safe context" in html.text
    docx = registry.extract(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        make_docx(),
        {},
    )
    assert "approval" in docx.text
    assert "owner | platform" in docx.text
    pdf = registry.extract("application/pdf", make_pdf(), {})
    assert pdf.metadata["pages"] == 1
    assert "citations" in pdf.text


def test_extractor_errors_and_registry_conflicts() -> None:
    registry = ExtractorRegistry()
    with pytest.raises(ValidationError, match="Invalid JSON"):
        registry.extract("application/json", b"{broken", {})
    with pytest.raises(ValidationError, match="Unsupported"):
        registry.extract("application/octet-stream", b"x", {})
    with pytest.raises(ValidationError, match="decoded"):
        registry.extract("text/plain", b"\xff", {"encoding": "utf-8"})
    with pytest.raises(ValueError, match="already registered"):
        registry.register(TextExtractor())
